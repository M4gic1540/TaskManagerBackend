"""Importa la planilla legacy de inventario (CSV o XLSX) directo por
consola. Uso:

    python manage.py import_legacy_inventory ruta/al/archivo.csv --actor-id 1 --actor-username admin

Hace lo mismo que POST /api/v1/inventory/import/, pero sin pasar por
HTTP: útil para la migración inicial masiva desde el servidor.

Recibe el actor por id+username en vez de resolverlo contra una tabla
`User` local: este servicio no tiene la tabla de usuarios en su propia
BD (microservicio con BD separada) — ver accounts/api para consultar
el id de un usuario existente."""
from django.core.management.base import BaseCommand

from accounts.enums import Role
from core.auth.jwt_claims_authentication import TokenClaimsUser
from inventory.import_mapping import row_to_asset_fields
from inventory.import_parser import parse_uploaded_inventory_file
from inventory.services.asset_service import AssetService


class _FileWrapper:
    """Envoltorio mínimo para reusar `parse_uploaded_inventory_file`
    (espera un objeto con `.name` y `.read()`, igual que un UploadedFile de DRF)."""

    def __init__(self, path):
        self.name = path
        self._path = path

    def read(self):
        with open(self._path, "rb") as f:
            return f.read()

    def load_workbook_source(self):
        return self._path


class Command(BaseCommand):
    help = "Importa el Excel/CSV legacy de inventario a la plataforma (reemplazo del intermediario Excel/GLPI)."

    def add_arguments(self, parser):
        parser.add_argument("path", type=str, help="Ruta al archivo .csv o .xlsx")
        parser.add_argument("--actor-id", type=int, required=True, help="Id del usuario que queda como autor de la importación")
        parser.add_argument("--actor-username", type=str, required=True, help="Username del usuario que queda como autor de la importación")

    def handle(self, *args, **options):
        path = options["path"]
        actor = TokenClaimsUser(
            id=options["actor_id"], username=options["actor_username"], role=Role.ADMIN
        )

        # openpyxl acepta directamente una ruta de archivo; csv.DictReader necesita
        # texto, así que para .csv seguimos usando el wrapper con .read().
        if path.lower().endswith(".xlsx"):
            import openpyxl

            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
            raw_rows = []
            for raw_row in rows_iter:
                if all(cell is None for cell in raw_row):
                    continue
                row = {h: ("" if c is None else str(c)) for h, c in zip(headers, raw_row) if h}
                raw_rows.append(row)
        else:
            raw_rows = parse_uploaded_inventory_file(_FileWrapper(path))

        if not raw_rows:
            self.stdout.write(self.style.WARNING("El archivo no tiene filas de datos."))
            return

        mapped_rows = [row_to_asset_fields(row) for row in raw_rows]
        result = AssetService().import_rows(rows=mapped_rows, actor=actor)

        self.stdout.write(self.style.SUCCESS(f"Creados: {result['created']}"))
        self.stdout.write(self.style.WARNING(f"Omitidos (ya importados): {result['skipped']}"))
        if result["errors"]:
            self.stdout.write(self.style.ERROR(f"Errores: {len(result['errors'])}"))
            for err in result["errors"]:
                self.stdout.write(f"  fila {err['row']}: {err['reason']}")
