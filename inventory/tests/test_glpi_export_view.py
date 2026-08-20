"""Test HTTP del export a Excel del inventario GLPI. No pega a MariaDB
real -mockea GLPIInventoryService.list_assets- solo verifica que la vista
arma un .xlsx válido con la cabecera esperada."""
import io
from unittest.mock import patch

import openpyxl
import pytest
from rest_framework.test import APIClient

from accounts.enums import Role
from core.auth.jwt_claims_authentication import TokenClaimsUser
from inventory.services.xlsx_export import GLPI_ASSETS_HEADERS

pytestmark = pytest.mark.django_db


@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture
def usuario_autenticado():
    return TokenClaimsUser(id=1, username="admin1", role=Role.ADMIN)


# PNG 1x1 válido, mínimo posible — alcanza para probar que se decodifica
# y embebe, no importa el contenido visual.
_TINY_PNG_B64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="

FAKE_ASSETS = [
    {
        "name": "PC Escritorio",
        "category": "PC",
        "brand": "HP",
        "serial": "SN123",
        "status": "En uso",
        "location": "CMM · Piso 3",
        "qr_code_data": f"data:image/png;base64,{_TINY_PNG_B64}",
    },
]


class TestGLPIAssetExportView:
    def test_export_devuelve_xlsx_con_la_misma_cabecera_de_la_tabla(self, api_client, usuario_autenticado):
        api_client.force_authenticate(user=usuario_autenticado)

        with patch(
            "inventory.api.views.GLPIInventoryService.list_assets",
            return_value=FAKE_ASSETS,
        ):
            response = api_client.get("/api/v1/inventory/glpi/export/")

        assert response.status_code == 200
        assert response["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "inventario_glpi.xlsx" in response["Content-Disposition"]

        wb = openpyxl.load_workbook(filename=io.BytesIO(response.content))
        ws = wb.active
        header_row = [cell.value for cell in ws[1]]
        assert header_row == GLPI_ASSETS_HEADERS

        data_row = [cell.value for cell in ws[2]]
        assert data_row[:6] == ["PC Escritorio", "PC", "HP", "SN123", "En uso", "CMM · Piso 3"]

        assert len(ws._images) == 1

    def test_export_requiere_autenticacion(self, api_client):
        response = api_client.get("/api/v1/inventory/glpi/export/")
        assert response.status_code == 401
