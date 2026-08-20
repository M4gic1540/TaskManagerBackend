"""Genera hojas de cálculo .xlsx a partir de listados de activos ya
resueltos (no hace consultas propias — recibe los dicts que arma
GLPIInventoryService.list_assets)."""
from __future__ import annotations

import base64
import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

GLPI_ASSETS_HEADERS = ["Nombre", "Categoría", "Marca", "Serial (SN)", "Estado", "Ubicación", "QR"]

_QR_COL_INDEX = len(GLPI_ASSETS_HEADERS)  # 1-based, última columna
_QR_CELL_PX = 70


def _qr_image(asset: Dict[str, Any]) -> XLImage | None:
    data_url = asset.get("qr_code_data")
    if not data_url or "," not in data_url:
        return None
    _, b64_data = data_url.split(",", 1)
    png_bytes = base64.b64decode(b64_data)
    img = XLImage(io.BytesIO(png_bytes))
    img.width = img.height = _QR_CELL_PX
    return img


def build_glpi_assets_xlsx(assets: List[Dict[str, Any]]) -> bytes:
    """Misma cabecera que la tabla de Inventario GLPI del frontend, con la
    columna QR al final como imagen embebida (Acción no va: es un botón,
    no dato)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario GLPI"
    ws.append(GLPI_ASSETS_HEADERS)

    for row_idx, asset in enumerate(assets, start=2):
        ws.append([
            asset.get("name") or "",
            asset.get("category") or "",
            asset.get("brand") or "",
            asset.get("serial") or "",
            asset.get("status") or "",
            asset.get("location") or "",
        ])
        ws.row_dimensions[row_idx].height = _QR_CELL_PX * 0.75

        img = _qr_image(asset)
        if img is not None:
            ws.add_image(img, f"{ws.cell(row=row_idx, column=_QR_COL_INDEX).coordinate}")

    for column_cells in ws.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 40)
    ws.column_dimensions[ws.cell(row=1, column=_QR_COL_INDEX).column_letter].width = 12

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
