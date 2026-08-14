"""Servicio de consulta en tiempo real a la base de datos MariaDB de GLPI.
Permite visualizar activos, filtrar, generar códigos QR de redirección directa a GLPI y KPIs."""
from __future__ import annotations

import base64
import io
import math
import zipfile
from typing import Any, Dict, List, Optional
import qrcode
from django.conf import settings
from django.db import connections
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from inventory.services.qr_label import build_qr_label_image, sanitize_filename
from core.resilience.circuit_breaker import call_with_breaker


TABLE_MAPPING = [
    ("glpi_computers", "PC", "computer.form.php"),
    ("glpi_monitors", "MONITOR", "monitor.form.php"),
    ("glpi_printers", "IMPRESORA", "printer.form.php"),
    ("glpi_networkequipments", "RED", "networkequipment.form.php"),
    ("glpi_peripherals", "PERIFERICO", "peripheral.form.php"),
    ("glpi_phones", "TELEFONO", "phone.form.php"),
]

# Los nombres de tabla no se pueden pasar como bind parameter (%s) en SQL:
# el driver solo parametriza valores, no identificadores. Se valida contra
# esta whitelist -derivada de TABLE_MAPPING, nunca de input externo- y se
# resuelve a un nombre ya quoteado antes de interpolarlo en la query, para
# que no dependa de que el llamador nunca le pase texto arbitrario.
_VALID_TABLE_NAMES = frozenset(name for name, _, _ in TABLE_MAPPING)


def _safe_table_identifier(connection, table_name: str) -> str:
    if table_name not in _VALID_TABLE_NAMES:
        raise ValueError(f"Tabla no permitida: {table_name!r}")
    return connection.ops.quote_name(table_name)


# Plantillas SQL estáticas: el token __TABLE__ se sustituye con
# str.replace() (nunca f-string/`.format()`/concatenación) por el
# identificador ya resuelto por _safe_table_identifier().
_LIST_ASSETS_SQL_TEMPLATE = """
    SELECT
        t.id,
        t.name,
        t.serial,
        t.otherserial AS legacy_id,
        t.comment,
        t.contact,
        m.name AS brand,
        st.name AS status,
        l.completename AS location,
        %s AS category,
        t.date_mod
    FROM __TABLE__ t
    LEFT JOIN glpi_manufacturers m ON t.manufacturers_id = m.id
    LEFT JOIN glpi_states st ON t.states_id = st.id
    LEFT JOIN glpi_locations l ON t.locations_id = l.id
    WHERE t.is_deleted = 0
"""

_CATEGORY_COUNT_SQL_TEMPLATE = "SELECT COUNT(*) FROM __TABLE__ WHERE is_deleted = 0"

_STATUS_COUNT_SQL_TEMPLATE = """
    SELECT COALESCE(st.name, 'Sin Estado') AS status_name, COUNT(*)
    FROM __TABLE__ t
    LEFT JOIN glpi_states st ON t.states_id = st.id
    WHERE t.is_deleted = 0
    GROUP BY st.name
"""

_LOCATION_COUNT_SQL_TEMPLATE = """
    SELECT COALESCE(l.completename, 'Sin Ubicación') AS loc_name, COUNT(*)
    FROM __TABLE__ t
    LEFT JOIN glpi_locations l ON t.locations_id = l.id
    WHERE t.is_deleted = 0 AND l.completename IS NOT NULL AND l.completename != ''
    GROUP BY l.completename
"""

_SERIAL_LOOKUP_SQL_TEMPLATE = """
    SELECT t.serial FROM __TABLE__ t
    WHERE t.id = %s AND t.is_deleted = 0
"""

_SERIAL_AND_INVENTORY_LOOKUP_SQL_TEMPLATE = """
    SELECT t.serial, t.otherserial FROM __TABLE__ t
    WHERE t.id = %s AND t.is_deleted = 0
"""


def _generate_qr_data_url(url: str) -> str:
    """Genera una imagen PNG del código QR codificado en Base64 Data URL, en
    alta resolución (mismos parámetros que _generate_qr_png_bytes) para que
    no se pixelee al imprimir la etiqueta física."""
    qr = qrcode.QRCode(box_size=8, border=2)
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    b64_str = base64.b64encode(buffer.getvalue()).decode("utf-8")
    return f"data:image/png;base64,{b64_str}"


class GLPIInventoryService:
    """Consulta la base de datos MariaDB (alias 'glpi' en settings.DATABASES)
    en modo solo lectura para la gestión y visualización del inventario."""

    def __init__(self, db_alias: str = "glpi"):
        self.db_alias = db_alias

    def _dictfetchall(self, cursor) -> List[Dict[str, Any]]:
        columns = [col[0] for col in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def list_assets(
        self,
        search: Optional[str] = None,
        category: Optional[str] = None,
        status: Optional[str] = None,
        limit: int = 200,
    ) -> List[Dict[str, Any]]:
        return call_with_breaker(self._list_assets_impl, search, category, status, limit)

    def _list_assets_impl(
        self,
        search: Optional[str] = None,
        category: Optional[str] = None,
        status: Optional[str] = None,
        limit: int = 200,
    ) -> List[Dict[str, Any]]:
        """Lista todos los activos desde MariaDB combinando las tablas de GLPI."""
        # No confiar en el type hint en runtime: se castea y acota antes de
        # usarlo como bind parameter, así un caller que pase algo no-numérico
        # falla temprano con un ValueError en vez de silenciarlo en SQL.
        safe_limit = max(1, min(int(limit), 1000))
        results: List[Dict[str, Any]] = []
        connection = connections[self.db_alias]
        glpi_base_url = getattr(settings, "GLPI_BASE_URL", "https://glpi.cmm.uchile.cl").rstrip("/")

        with connection.cursor() as cursor:
            for table_name, cat_code, form_page in TABLE_MAPPING:
                if category and category.upper() not in ("ALL", cat_code):
                    continue

                # Query estática con placeholder de token (no f-string): el
                # nombre de tabla ya validado/quoteado se sustituye vía
                # .replace(), separado de la interpolación de valores (%s).
                sql = _LIST_ASSETS_SQL_TEMPLATE.replace(
                    "__TABLE__", _safe_table_identifier(connection, table_name)
                )
                params: list = [cat_code]

                if status:
                    sql += " AND LOWER(st.name) LIKE %s"
                    params.append(f"%{status.lower()}%")

                if search:
                    sql += """ AND (
                        LOWER(t.name) LIKE %s OR 
                        LOWER(t.serial) LIKE %s OR 
                        LOWER(t.otherserial) LIKE %s OR 
                        LOWER(t.contact) LIKE %s OR 
                        LOWER(l.completename) LIKE %s OR 
                        LOWER(m.name) LIKE %s
                    )"""
                    search_pattern = f"%{search.lower()}%"
                    params.extend([search_pattern] * 6)

                sql += " ORDER BY t.id DESC LIMIT %s"
                params.append(safe_limit)

                cursor.execute(sql, params)
                rows = self._dictfetchall(cursor)
                for r in rows:
                    r["id_str"] = f"{cat_code.lower()}_{r['id']}"
                    r["code"] = f"GLPI-{cat_code}-{r['id']}"
                    # legacy_id viene de t.otherserial: el número de inventario
                    # real cargado en GLPI (no `code`, que es sintético). Muchos
                    # activos no lo tienen cargado, por eso puede quedar None.
                    r["inventory_number"] = (r.get("legacy_id") or "").strip() or None
                    r["glpi_url"] = f"{glpi_base_url}/front/{form_page}?id={r['id']}"
                    r["qr_code_data"] = _generate_qr_data_url(r["glpi_url"])

                results.extend(rows)

        # Ordenar por fecha de modificación / id
        results.sort(key=lambda x: str(x.get("date_mod") or ""), reverse=True)
        return results[:limit]

    def get_dashboard_summary(self) -> Dict[str, Any]:
        return call_with_breaker(self._get_dashboard_summary_impl)

    def _get_dashboard_summary_impl(self) -> Dict[str, Any]:
        """Calcula los KPIs agregados del inventario consultando MariaDB."""
        connection = connections[self.db_alias]
        by_category: Dict[str, int] = {}
        by_status: Dict[str, int] = {}
        top_locations: Dict[str, int] = {}
        total = 0

        with connection.cursor() as cursor:
            for table_name, cat_code, _ in TABLE_MAPPING:
                safe_table = _safe_table_identifier(connection, table_name)

                # Total por categoría
                cursor.execute(_CATEGORY_COUNT_SQL_TEMPLATE.replace("__TABLE__", safe_table))
                count = cursor.fetchone()[0]
                by_category[cat_code] = count
                total += count

                # Estados
                cursor.execute(_STATUS_COUNT_SQL_TEMPLATE.replace("__TABLE__", safe_table))
                for st_name, st_count in cursor.fetchall():
                    by_status[st_name] = by_status.get(st_name, 0) + st_count

                # Ubicaciones top
                cursor.execute(_LOCATION_COUNT_SQL_TEMPLATE.replace("__TABLE__", safe_table))
                for loc_name, loc_count in cursor.fetchall():
                    top_locations[loc_name] = top_locations.get(loc_name, 0) + loc_count

        sorted_locations = dict(sorted(top_locations.items(), key=lambda item: item[1], reverse=True)[:10])

        return {
            "total": total,
            "by_category": by_category,
            "by_status": by_status,
            "top_locations": sorted_locations,
        }

    def _generate_qr_png_bytes(self, url: str, *, box_size: int = 8) -> bytes:
        """Genera QR en alta resolución para impresión."""
        qr = qrcode.QRCode(box_size=box_size, border=2)
        qr.add_data(url)
        qr.make(fit=True)
        img = qr.make_image(fill_color="black", back_color="white")
        buffer = io.BytesIO()
        img.save(buffer, format="PNG")
        return buffer.getvalue()

    def generate_qr_sheet_docx(self, *, asset_ids: list[str]) -> bytes:
        return call_with_breaker(self._generate_qr_sheet_docx_impl, asset_ids=asset_ids)

    def _generate_qr_sheet_docx_impl(self, *, asset_ids: list[str]) -> bytes:
        """Genera documento Word con QRs de GLPI y S/N debajo. asset_ids es lista
        de strings formato 'pc_123' o 'monitor_45' (category_id). Solo muestra QR + S/N."""
        connection = connections[self.db_alias]
        glpi_base_url = getattr(settings, "GLPI_BASE_URL", "https://glpi.cmm.uchile.cl").rstrip("/")

        assets_data = []
        for asset_id in asset_ids:
            if "_" not in asset_id:
                continue
            cat_code, obj_id = asset_id.split("_", 1)
            obj_id = int(obj_id)

            for table_name, category, form_page in TABLE_MAPPING:
                if cat_code.upper() != category:
                    continue

                with connection.cursor() as cursor:
                    safe_table = _safe_table_identifier(connection, table_name)
                    cursor.execute(
                        _SERIAL_LOOKUP_SQL_TEMPLATE.replace("__TABLE__", safe_table),
                        [obj_id],
                    )
                    row = cursor.fetchone()
                    if row:
                        glpi_url = f"{glpi_base_url}/front/{form_page}?id={obj_id}"
                        assets_data.append({
                            "serial": row[0] or "",
                            "glpi_url": glpi_url,
                        })
                break

        if not assets_data:
            from core.exceptions import EntityNotFoundError
            raise EntityNotFoundError("Ninguno de los activos solicitados existe.")

        document = Document()
        document.add_heading("Etiquetas QR de Inventario GLPI", level=1)

        columns = 3
        rows = math.ceil(len(assets_data) / columns)
        table = document.add_table(rows=rows, cols=columns)

        for index, asset in enumerate(assets_data):
            cell = table.cell(index // columns, index % columns)

            image_paragraph = cell.paragraphs[0]
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_paragraph.add_run().add_picture(
                io.BytesIO(self._generate_qr_png_bytes(asset["glpi_url"])), width=Cm(2.5)
            )

            if asset["serial"]:
                serial_paragraph = cell.add_paragraph()
                serial_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                serial_paragraph.add_run(asset["serial"]).bold = True

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def generate_qr_images_zip(self, *, asset_ids: list[str]) -> bytes:
        return call_with_breaker(self._generate_qr_images_zip_impl, asset_ids=asset_ids)

    def _generate_qr_images_zip_impl(self, *, asset_ids: list[str]) -> bytes:
        """Genera un ZIP con imágenes PNG de QRs (QR + S/N).
        asset_ids es lista de strings formato 'pc_123' o 'monitor_45'."""
        connection = connections[self.db_alias]
        glpi_base_url = getattr(settings, "GLPI_BASE_URL", "https://glpi.cmm.uchile.cl").rstrip("/")

        assets_data = []
        for asset_id in asset_ids:
            if "_" not in asset_id:
                continue
            cat_code, obj_id = asset_id.split("_", 1)
            obj_id = int(obj_id)

            for table_name, category, form_page in TABLE_MAPPING:
                if cat_code.upper() != category:
                    continue

                with connection.cursor() as cursor:
                    safe_table = _safe_table_identifier(connection, table_name)
                    cursor.execute(
                        _SERIAL_AND_INVENTORY_LOOKUP_SQL_TEMPLATE.replace("__TABLE__", safe_table),
                        [obj_id],
                    )
                    row = cursor.fetchone()
                    if row:
                        glpi_url = f"{glpi_base_url}/front/{form_page}?id={obj_id}"
                        assets_data.append({
                            "serial": row[0] or "SIN_SN",
                            "inventory_number": (row[1] or "").strip() or None,
                            "glpi_url": glpi_url,
                        })
                break

        if not assets_data:
            from core.exceptions import EntityNotFoundError
            raise EntityNotFoundError("Ninguno de los activos solicitados existe.")

        # Crear ZIP
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            used_names = set()
            for asset in assets_data:
                filename = f"{sanitize_filename(asset['serial'])}.png"

                # Si el nombre ya existe, agregar sufijo
                if filename in used_names:
                    base, ext = filename.rsplit(".", 1)
                    suffix = 2
                    while f"{base}_{suffix}.{ext}" in used_names:
                        suffix += 1
                    filename = f"{base}_{suffix}.{ext}"

                used_names.add(filename)
                image_buffer = build_qr_label_image(
                    target_url=asset["glpi_url"],
                    serial_number=asset["serial"],
                    inventory_number=asset["inventory_number"],
                )
                zip_file.writestr(filename, image_buffer.getvalue())

        zip_buffer.seek(0)
        return zip_buffer.getvalue()
